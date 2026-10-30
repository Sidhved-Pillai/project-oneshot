from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from src.trip_dtr_report import DTR_REVIEW_COLUMNS, export_operational_dtr
from src.ai_intake import DTR_FIELD_MAP, DTRIntakeRow, UnifiedIntakeRow, UnifiedIntakeResult
from src.invoice_recovery import recover_entry_fields
from tests.test_batched_interactions import trip_app


def test_dtr_column_order_and_export_preserve_driver_and_diesel_rate():
    columns = DTR_REVIEW_COLUMNS
    assert not set(['Company Freight','Bill No.','Loading & Unloading','SG & Bisleri Damages','Debit Amt.','Review Notes']) & set(columns)
    assert columns[columns.index('Date')+1:columns.index('Vehicle No.')+1] == ['Driver Name','Vehicle No.']
    start = columns.index('UPI')
    assert columns[start:start+5] == ['UPI','Diesel Pump Name','Diesel Qty','Diesel Rate','Diesel Adv.']
    sheet = load_workbook(BytesIO(export_operational_dtr(pd.DataFrame([{'Driver Name':'ANSARI','Diesel Rate':95.5}]))))['DTR']
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2,headers.index('Driver Name')+1).value == 'ANSARI'
    assert sheet.cell(2,headers.index('Diesel Rate')+1).value == 95.5
    assert DTR_FIELD_MAP['Driver Name'] == 'driver_name'
    assert DTR_FIELD_MAP['Diesel Rate'] == 'diesel_rate'
    assert set(DTR_FIELD_MAP.values()) <= set(DTRIntakeRow.model_fields)


def test_saint_gobain_invoice_table_recovery_keeps_one_lr():
    first = UnifiedIntakeRow(company_name='Saint-Gobain India Private Limited', date='2026-10-20',
        branch='Pune', from_location='Bhiwandi', to_location='Pune', vehicle_number='MH05EL1923',
        lr_number='00000010036848-A', invoice_number='9102085603', invoice_numbers=['9102085603'])
    calls = []
    def extractor(*args, **kwargs):
        calls.append(kwargs)
        return UnifiedIntakeResult(rows=[first.model_copy(update={'invoice_numbers':['9102085603','9102085604']})]), None
    result = recover_entry_fields(UnifiedIntakeResult(rows=[first]), extractor, 'test','context',
        [{'mime_type':'image/jpeg','data':b'example'}], None, ['Pune'], [])
    assert len(result.rows) == 1 and len(calls) == 1 and calls[0]['high_detail']
    assert result.rows[0].invoice_numbers == ['9102085603','9102085604']
    assert result.rows[0].lr_number == '00000010036848-A'


def test_saint_gobain_entry_has_separate_invoices_and_single_lr():
    app = trip_app('Nitish')
    next(w for w in app.text_input if w.label == 'Company name *').set_value('Saint-Gobain India Private Limited').run()
    assert not app.exception
    next(w for w in app.text_input if w.label == 'Invoice No. 1 *').set_value('9102085603')
    next(w for w in app.text_input if w.label == 'Invoice No. 2').set_value('9102085604')
    next(w for w in app.text_input if w.label == 'LR number').set_value('00000010036848-A').run()
    assert not app.exception
    values = app.session_state['test_returned_values']
    assert '9102085603' in values['invoice_number'] and '9102085604' in values['invoice_number']
    assert values['lr_number'] == '00000010036848-A'
