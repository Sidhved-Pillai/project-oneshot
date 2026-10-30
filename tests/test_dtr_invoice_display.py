from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from src.invoice_display import dtr_invoice_text
from src.invoice_numbers import invoice_identifier_tokens, user_invoice_duplicates
from src.trip_dtr_report import export_operational_dtr


def test_dtr_lists_all_invoices_and_preserves_structured_ids():
    assert dtr_invoice_text('test123 / test1234') == 'test123, test1234'
    assert dtr_invoice_text(['test1','test2','test3','test4','test5','test6']) == 'test1, test2, test3, test4, test5, test6'
    assert dtr_invoice_text('GST/GST/0885/26-27 / GST/GST/1087/26-27') == 'GST/GST/0885/26-27, GST/GST/1087/26-27'
    assert dtr_invoice_text('test123, test1234') == 'test123, test1234'
    workbook = load_workbook(BytesIO(export_operational_dtr(pd.DataFrame([{'Invoice No.':'test123 / test1234'}]))))
    sheet = workbook['DTR']
    column = [cell.value for cell in sheet[1]].index('Invoice No.') + 1
    assert sheet.cell(2,column).value == 'test123, test1234'
    assert sheet.cell(2,column).number_format == '@'


def test_comma_separated_saved_invoices_remain_individually_duplicate_checked():
    assert invoice_identifier_tokens('test123, test1234') == {'test123','test1234'}
    row = {'created_by':'Nitish','report_scope':'Both','invoice_number':'test123, test1234'}
    assert user_invoice_duplicates([row],'Nitish','test1234') == [row]
    assert not user_invoice_duplicates([row],'Nitish','test12345')
