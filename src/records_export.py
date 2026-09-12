"""Complete, human-readable export of filtered Records-tab rows."""

import json
from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pnl_report import DIRECT_EXPENSE_COLUMNS


EXPENSE_EXPORT_COLUMNS = [*DIRECT_EXPENSE_COLUMNS, "Passing expense"]
RECORD_EXPORT_COLUMNS = [
    "Record", "Record Type", "Date", "Status", "Created By", "Branch", "Company Name",
    "Vehicle Number", "Vehicle Capacity", "Own / Outside", "From", "To", "LR Number",
    "Invoice Number", "Beneficiary Name", "Account Number", "IFSC Code", "Transporter Name",
    "Vehicle Placed By", "Revenue Freight", "Transporter Freight", "RTGS", "Cash", "UPI",
    "Diesel", "Diesel Qty", "Diesel Rate", "Toll Expense", "Repairs & Maintenance",
    "Repair Reason", "Billtee", "Total Advance", "Balance Amount", "Payment", "Payment Mode",
    "Diesel Pump Name", "Card Name", "Remarks", "Invoice Evidence Filename",
    *EXPENSE_EXPORT_COLUMNS, "Insurance Period Start", "Insurance Period End",
    "Vehicle Tax Period Start", "Vehicle Tax Period End", "Card", "Cardholders Name",
]


def _json(value):
    if isinstance(value, dict):
        return value
    try:
        return json.loads(value or "{}")
    except (TypeError, json.JSONDecodeError):
        return {}


def _period_bounds(value):
    if isinstance(value, dict):
        return value.get("start", ""), value.get("end", "")
    if isinstance(value, (list, tuple)):
        return (value[0] if len(value) > 0 else "", value[1] if len(value) > 1 else "")
    return "", ""


def records_export_rows(records):
    output = []
    for row in records:
        dtr, rtgs = _json(row.get("dtr_data")), _json(row.get("rtgs_data"))
        categories = dtr.get("categories", {}) if isinstance(dtr.get("categories"), dict) else {}
        periods = dtr.get("periods", {}) if isinstance(dtr.get("periods"), dict) else {}
        insurance_start, insurance_end = _period_bounds(periods.get("Insurance"))
        tax_start, tax_end = _period_bounds(periods.get("Vehicle Tax"))
        payments = dtr.get("payments", {}) if isinstance(dtr.get("payments"), dict) else {}
        item = {
            "Record": row.get("request_number", ""),
            "Record Type": "Direct Expense" if row.get("report_scope") == "Expense" else "Trip",
            "Date": row.get("trip_date"), "Status": row.get("status", ""),
            "Created By": row.get("created_by", ""), "Branch": row.get("branch", ""),
            "Company Name": row.get("company_name", ""), "Vehicle Number": row.get("vehicle_number", ""),
            "Vehicle Capacity": row.get("vehicle_type", ""), "Own / Outside": row.get("ownership_type", ""),
            "From": row.get("from_location", ""), "To": row.get("to_location", ""),
            "LR Number": dtr.get("LR No.", ""), "Invoice Number": dtr.get("Invoice No.") or row.get("invoice_number", ""),
            "Beneficiary Name": row.get("beneficiary_name", ""), "Account Number": rtgs.get("BENE_ACC_NO", ""),
            "IFSC Code": rtgs.get("BENE_IFSC", ""), "Transporter Name": row.get("transporter_name", ""),
            "Vehicle Placed By": dtr.get("Veh Placed by", ""), "Revenue Freight": row.get("revenue", 0),
            "Transporter Freight": row.get("transporter_freight", 0), "RTGS": row.get("rtgs_advance", 0),
            "Cash": row.get("cash_advance", 0), "UPI": row.get("upi", 0),
            "Diesel": row.get("diesel_advance", 0), "Diesel Qty": row.get("diesel_quantity", 0),
            "Diesel Rate": dtr.get("Diesel Rate", 0), "Toll Expense": dtr.get("Toll Expense", 0),
            "Repairs & Maintenance": dtr.get("Repairs & Maintenance", 0),
            "Repair Reason": dtr.get("Repair Reason", ""), "Billtee": dtr.get("Billtee", 0),
            "Total Advance": row.get("total_advance", 0), "Balance Amount": row.get("balance_amount", 0),
            "Payment": row.get("payment", 0), "Payment Mode": row.get("payment_mode", ""),
            "Diesel Pump Name": dtr.get("Diesel Pump Name", ""), "Card Name": dtr.get("Card Name", ""),
            "Remarks": row.get("notes", ""), "Invoice Evidence Filename": row.get("source_filename", ""),
            "Insurance Period Start": insurance_start, "Insurance Period End": insurance_end,
            "Vehicle Tax Period Start": tax_start, "Vehicle Tax Period End": tax_end,
            "Card": payments.get("Card", 0), "Cardholders Name": dtr.get("Cardholders Name", ""),
        }
        item.update({column: categories.get(column, 0) for column in EXPENSE_EXPORT_COLUMNS})
        output.append({column: item.get(column, "") for column in RECORD_EXPORT_COLUMNS})
    return output


def export_records_excel(records):
    frame = pd.DataFrame(records_export_rows(records), columns=RECORD_EXPORT_COLUMNS)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl", date_format="DD-MM-YYYY") as writer:
        frame.to_excel(writer, index=False, sheet_name="Records")
        sheet = writer.book["Records"]
        fill = PatternFill("solid", fgColor="0F766E")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        money_columns = set(RECORD_EXPORT_COLUMNS[19:37]) | set(EXPENSE_EXPORT_COLUMNS) | {"Card"}
        for index, name in enumerate(RECORD_EXPORT_COLUMNS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(32, max(13, len(name) + 2))
            if name in money_columns:
                for row_index in range(2, sheet.max_row + 1):
                    sheet.cell(row=row_index, column=index).number_format = '#,##0.00'
    return output.getvalue()
